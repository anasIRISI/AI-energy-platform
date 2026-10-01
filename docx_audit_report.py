from __future__ import annotations

import json
import re
import sys
import zipfile
from pathlib import Path

from docx import Document


def main(path: str) -> None:
    document = Document(path)
    paragraphs = []
    headings = []
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        style = paragraph.style.name if paragraph.style else ""
        if text:
            item = {"index": index, "style": style, "text": text}
            paragraphs.append(item)
            if style.startswith("Heading"):
                headings.append(item)

    ai_markers = re.compile(r"\b(IA|intelligence artificielle|chatbot|RAG|FastAPI|Gemini|Chroma|simulation|worker)\b", re.I)
    ai_paragraphs = [item for item in paragraphs if ai_markers.search(item["text"])]

    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml").decode("utf-8", errors="replace")
        settings = archive.read("word/settings.xml").decode("utf-8", errors="replace") if "word/settings.xml" in archive.namelist() else ""

    report = {
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "inline_shape_count": len(document.inline_shapes),
        "section_count": len(document.sections),
        "heading_count": len(headings),
        "headings": headings,
        "ai_paragraphs": ai_paragraphs,
        "has_toc_field": "TOC" in xml,
        "has_figure_field": "SEQ Figure" in xml,
        "has_table_field": "SEQ Table" in xml,
        "updates_fields_on_open": "updateFields" in settings,
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main(sys.argv[1])
