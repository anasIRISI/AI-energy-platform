from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"C:\Users\hp\Documents\AI-energy-platform")
SOURCE = ROOT / "Rapport_PFA_EcoReno_Chapitre_IA_Detaille.docx"
OUTPUT = ROOT / "Rapport_PFA_EcoReno_Final_IA_Detaillee.docx"
FONT = "Times New Roman"
BLACK = RGBColor(0, 0, 0)


def set_run_font(run, size=12, bold=False, italic=False):
    run.font.name = FONT
    if run._element.rPr is None:
        run._element.get_or_add_rPr()
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = BLACK


def set_style_font(style, size, bold=False, italic=False):
    style.font.name = FONT
    style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = BLACK
    style.paragraph_format.space_before = Pt(12 if size >= 14 else 8)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.keep_with_next = True


def format_body(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.15
    for run in paragraph.runs:
        set_run_font(run, 12)


def format_heading(paragraph, level):
    paragraph.style = f"Heading {level}"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(16 if level == 1 else 11)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.page_break_before = level == 1
    for run in paragraph.runs:
        set_run_font(run, 16 if level == 1 else 13, bold=True)


def add_complex_field(run, instruction: str, result: str = "0"):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = result
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, placeholder, end])


def add_simple_field_after(paragraph, instruction: str):
    field_paragraph = OxmlElement("w:p")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), instruction)
    field_paragraph.append(field)
    paragraph._p.addnext(field_paragraph)


def remove_paragraph(paragraph):
    paragraph._element.getparent().remove(paragraph._element)


def remove_table(table):
    table._element.getparent().remove(table._element)


def insert_paragraph_after(paragraph, text="", style=None):
    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    created = Paragraph(element, paragraph._parent)
    if style:
        created.style = style
    if text:
        created.add_run(text)
    return created


def add_requirements_table(document, after_paragraph):
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ("Catégorie", "Exigence", "Priorité")
    for cell, text in zip(table.rows[0].cells, headers):
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "1A6D45")
        cell._tc.get_or_add_tcPr().append(shading)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cell.paragraphs[0].add_run(text)
        set_run_font(run, 10, bold=True)
        run.font.color.rgb = RGBColor(255, 255, 255)
    rows = (
        ("Information", "Consulter les solutions, produits et contenus énergétiques.", "Élevée"),
        ("Simulation", "Saisir les données du logement et obtenir un résultat sauvegardé.", "Élevée"),
        ("Assistant", "Poser une question, recevoir une réponse contextualisée et être guidé.", "Élevée"),
        ("Rendez-vous", "Créer, modifier ou annuler un rendez-vous après confirmation.", "Moyenne"),
        ("Administration", "Gérer les produits, régions et informations de la plateforme.", "Élevée"),
    )
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            if index % 2:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "F4FAF6")
                cell._tc.get_or_add_tcPr().append(shading)
            run = cell.paragraphs[0].add_run(text)
            set_run_font(run, 9)
            cell.paragraphs[0].paragraph_format.space_after = Pt(2)
    after_paragraph._p.addnext(table._tbl)
    return table


def replace_caption(paragraph, number: int):
    text = paragraph.text.strip()
    description = re.sub(r"^Figure\s*(?::|\d+)?\s*", "", text, flags=re.IGNORECASE).strip(" :")
    paragraph.clear()
    paragraph.style = "Caption"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run("Figure ")
    set_run_font(run, 10, italic=True)
    add_complex_field(run, "SEQ Figure \\* ARABIC", str(number))
    ending = paragraph.add_run(f" {description}")
    set_run_font(ending, 10, italic=True)


def replace_table_caption(paragraph, number: int):
    text = paragraph.text.strip()
    description = re.sub(r"^Table\s*(?::|\d+)?\s*", "", text, flags=re.IGNORECASE).strip(" :")
    paragraph.clear()
    paragraph.style = "Caption"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(7)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run("Tableau ")
    set_run_font(run, 10, italic=True)
    add_complex_field(run, "SEQ Table \\* ARABIC", str(number))
    ending = paragraph.add_run(f" {description}")
    set_run_font(ending, 10, italic=True)


def add_footer_page_number(section):
    section.different_first_page_header_footer = True
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    run = paragraph.add_run("Page ")
    set_run_font(run, 10)
    add_complex_field(run, "PAGE")


def ensure_update_fields(document):
    settings = document.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def main():
    document = Document(SOURCE)

    if "Caption" not in [style.name for style in document.styles]:
        document.styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(document.styles["Heading 1"], 16, bold=True)
    set_style_font(document.styles["Heading 2"], 13, bold=True)
    set_style_font(document.styles["Heading 3"], 12, bold=True)
    set_style_font(document.styles["Caption"], 10, italic=True)

    paragraphs = list(document.paragraphs)
    cover_end = next(index for index, paragraph in enumerate(paragraphs) if paragraph.text.strip() == "Remerciement")

    # The retained report contains many empty spacer paragraphs, including
    # page-break-only paragraphs. They create blank pages once Heading 1 is
    # configured to begin a page. Remove only body spacers after the cover;
    # visible images, section properties and actual content remain untouched.
    for paragraph in list(document.paragraphs[cover_end + 1:]):
        xml = paragraph._p.xml
        if (not paragraph.text.strip() or paragraph.text.strip() == ".") and "<w:drawing" not in xml and "<w:sectPr" not in xml:
            remove_paragraph(paragraph)

    main_titles = {
        "Remerciement", "Résumé", "Introduction Générale",
        "Conclusion générale",
    }
    main_title_fragments = (
        "Présentation de l’entreprise", "Présentation du projet",
        "Conception et Architecture", "Réalisation du projet",
    )

    caption_number = 1
    table_number = 1
    for index, paragraph in enumerate(list(document.paragraphs)):
        text = paragraph.text.strip()
        if not text:
            continue
        if index < cover_end:
            continue
        if text in main_titles or any(fragment in text for fragment in main_title_fragments):
            format_heading(paragraph, 1)
        elif paragraph.style.name == "style2" or re.match(r"^4\.\d+(?:\.\d+)?\s", text):
            format_heading(paragraph, 2)
            if text.startswith("4 Réalisation du module"):
                paragraph.paragraph_format.page_break_before = True
        elif text.lower().startswith("figure") and "liste des figures" not in text.lower():
            replace_caption(paragraph, caption_number)
            caption_number += 1
        elif text.lower().startswith("table ") and "table des matières" not in text.lower():
            replace_table_caption(paragraph, table_number)
            table_number += 1
        elif paragraph.style.name == "List Paragraph":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.line_spacing = 1.1
            for run in paragraph.runs:
                set_run_font(run, 12)
        else:
            format_body(paragraph)

    # Expand the front-matter pages so their academic layout is balanced while
    # keeping the cover and the original section order unchanged.
    resume_heading = next(paragraph for paragraph in document.paragraphs if paragraph.text.strip() == "Résumé")
    p = resume_heading.insert_paragraph_before("")
    p.add_run("Je souhaite également remercier mes collègues et toutes les personnes qui ont contribué, directement ou indirectement, à l'avancement de ce travail. Leurs échanges, leurs remarques et leur disponibilité ont constitué un appui précieux durant les différentes étapes de conception, de développement et de validation de l'application.")
    format_body(p)
    p = resume_heading.insert_paragraph_before("")
    p.add_run("Enfin, je remercie ma famille et mes proches pour leur confiance, leur patience et leurs encouragements. Leur soutien a été essentiel pour mener ce projet à son terme dans de bonnes conditions.")
    format_body(p)

    toc_title_for_summary = next(paragraph for paragraph in document.paragraphs if paragraph.text.strip() in {"Tables de matières", "Table des matières"})
    p = toc_title_for_summary.insert_paragraph_before("")
    p.add_run("La solution développée s'appuie sur une architecture web moderne associant React pour l'interface, Spring Boot pour la gestion métier, PostgreSQL pour la persistance des données et FastAPI pour les services d'intelligence artificielle. Elle intègre une base documentaire, un chatbot et un module de simulation énergétique afin de proposer un accompagnement personnalisé.")
    format_body(p)
    p = toc_title_for_summary.insert_paragraph_before("")
    p.add_run("Le rapport présente le contexte et les objectifs du projet, les choix d'architecture, les principales étapes de réalisation ainsi que les résultats obtenus. Une attention particulière est portée au module d'intelligence artificielle, notamment au chatbot Luna, à la recherche documentaire RAG et au fonctionnement de la simulation énergétique.")
    format_body(p)

    # Dynamic lists are inserted after their existing report titles.
    toc_title = next(paragraph for paragraph in document.paragraphs if paragraph.text.strip() == "Tables de matières")
    toc_title.text = "Table des matières"
    toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    toc_title.paragraph_format.space_after = Pt(10)
    toc_title.paragraph_format.page_break_before = True
    for run in toc_title.runs:
        set_run_font(run, 16, bold=True)
    add_simple_field_after(toc_title, 'TOC \\o "1-3" \\h \\z \\u')

    figure_list_title = next(paragraph for paragraph in document.paragraphs if paragraph.text.strip() == "Liste des figures")
    table_list_title = figure_list_title.insert_paragraph_before("Liste des tableaux", style="Normal")
    table_list_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table_list_title.paragraph_format.space_after = Pt(10)
    for run in table_list_title.runs:
        set_run_font(run, 16, bold=True)
    add_simple_field_after(table_list_title, 'TOC \\h \\z \\c "Table"')

    figure_list_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure_list_title.paragraph_format.space_after = Pt(10)
    for run in figure_list_title.runs:
        set_run_font(run, 16, bold=True)
    add_simple_field_after(figure_list_title, 'TOC \\h \\z \\c "Figure"')

    conclusion = document.add_paragraph("Conclusion générale", style="Heading 1")
    conclusion.paragraph_format.page_break_before = True
    format_heading(conclusion, 1)
    body = document.add_paragraph(style="Normal")
    body.add_run("Ce rapport a présenté la conception et la réalisation de la plateforme EcoReno+, depuis l'architecture générale jusqu'aux interfaces, au backend et au module IA. La solution repose sur une séparation claire entre React, Spring Boot, PostgreSQL et FastAPI afin de faciliter l'évolution de chaque composant.")
    format_body(body)
    body = document.add_paragraph(style="Normal")
    body.add_run("Les fonctionnalités réalisées permettent de consulter les solutions énergétiques, de collecter les informations nécessaires à une simulation et de centraliser les résultats. Le module IA intègre le chatbot Luna connecté au backend, une recherche documentaire régionale et un calcul explicable des estimations. L'amélioration attendue concerne surtout l'enrichissement du corpus et l'évaluation de modèles supervisés sur des données de référence.")
    format_body(body)
    body = document.add_paragraph(style="Normal")
    body.add_run("Cette démarche fournit une base évolutive pour accompagner les utilisateurs dans leurs projets de rénovation énergétique, tout en préservant la traçabilité des sources et la prudence nécessaire aux conseils et estimations présentés.")
    format_body(body)

    for section in document.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(1.7)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        add_footer_page_number(section)

    ensure_update_fields(document)
    document.core_properties.title = "Rapport PFA EcoReno"
    document.core_properties.subject = "Rapport final avec module Intelligence Artificielle"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
