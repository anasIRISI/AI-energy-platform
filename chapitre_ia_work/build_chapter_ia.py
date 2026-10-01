from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"C:\Users\hp\Documents\AI-energy-platform")
WORK = ROOT / "chapitre_ia_work"
SOURCE = Path(r"C:\Users\hp\Downloads\Page de garde.docx")
OUTPUT = ROOT / "Rapport_PFA_EcoReno_Chapitre_IA_Detaille.docx"

FONT = "Times New Roman"
BLACK = RGBColor(0, 0, 0)
GREEN = (20, 109, 69)
ORANGE = (232, 144, 32)
PALE_GREEN = (235, 247, 239)
PALE_ORANGE = (255, 246, 228)
GRAY = (88, 96, 92)


def font(size: int, bold: bool = False):
    candidates = [
        r"C:\Windows\Fonts\timesbd.ttf" if bold else r"C:\Windows\Fonts\times.ttf",
        r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def wrap(draw: ImageDraw.ImageDraw, text: str, max_width: int, active_font):
    words, lines, current = text.split(), [], ""
    for word in words:
        proposal = f"{current} {word}".strip()
        if current and draw.textbbox((0, 0), proposal, font=active_font)[2] > max_width:
            lines.append(current)
            current = word
        else:
            current = proposal
    if current:
        lines.append(current)
    return lines


def draw_centered(draw, box, text, active_font, fill=(0, 0, 0), spacing=6):
    x1, y1, x2, y2 = box
    lines = wrap(draw, text, x2 - x1 - 30, active_font)
    heights = [draw.textbbox((0, 0), line, font=active_font)[3] for line in lines]
    height = sum(heights) + max(0, len(lines) - 1) * spacing
    y = y1 + ((y2 - y1) - height) / 2
    for line, line_height in zip(lines, heights):
        width = draw.textbbox((0, 0), line, font=active_font)[2]
        draw.text((x1 + ((x2 - x1) - width) / 2, y), line, font=active_font, fill=fill)
        y += line_height + spacing


def rounded_box(draw, box, text, fill, outline, text_fill=(0, 0, 0), radius=22, size=34):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=3)
    draw_centered(draw, box, text, font(size, True), text_fill)


def arrow(draw, start, end, fill=GREEN, width=5):
    draw.line([start, end], fill=fill, width=width)
    dx, dy = end[0] - start[0], end[1] - start[1]
    length = max((dx * dx + dy * dy) ** 0.5, 1)
    ux, uy = dx / length, dy / length
    px, py = -uy, ux
    tip = end
    left = (end[0] - 20 * ux + 10 * px, end[1] - 20 * uy + 10 * py)
    right = (end[0] - 20 * ux - 10 * px, end[1] - 20 * uy - 10 * py)
    draw.polygon([tip, left, right], fill=fill)


def architecture_figure(path: Path):
    image = Image.new("RGB", (1800, 780), "white")
    draw = ImageDraw.Draw(image)
    draw.text((80, 42), "Architecture du module IA observée dans le dépôt", font=font(44, True), fill=(0, 0, 0))
    rounded_box(draw, (90, 210, 430, 420), "Frontend React\nFormulaire, catalogue et interface", PALE_GREEN, GREEN)
    rounded_box(draw, (560, 210, 920, 420), "Backend Spring Boot\nDonnées métier et orchestration", PALE_ORANGE, ORANGE)
    rounded_box(draw, (1050, 105, 1690, 295), "Chatbot Luna FastAPI :8081\nGemini 2.0 Flash et actions métier", PALE_GREEN, GREEN)
    rounded_box(draw, (1050, 390, 1335, 665), "Service IA FastAPI :8001\nRAG, simulations et worker", PALE_ORANGE, ORANGE)
    rounded_box(draw, (1410, 390, 1690, 525), "ChromaDB\nBase vectorielle", (242, 247, 255), (72, 119, 180))
    rounded_box(draw, (1410, 565, 1690, 735), "Gemini\nConseils de simulation\nsi configuré", PALE_GREEN, GREEN)
    arrow(draw, (430, 315), (560, 315))
    arrow(draw, (920, 315), (1050, 200))
    arrow(draw, (920, 335), (1050, 525))
    arrow(draw, (1335, 525), (1410, 460))
    arrow(draw, (1335, 565), (1410, 650))
    draw.text((95, 745), "Échanges REST entre services ; les simulations et les conversations sont enregistrées par Spring Boot.", font=font(26), fill=GRAY)
    image.save(path)


def chatbot_figure(path: Path):
    image = Image.new("RGB", (1800, 760), "white")
    draw = ImageDraw.Draw(image)
    draw.text((80, 42), "Fonctionnement du chatbot Luna", font=font(44, True), fill=(0, 0, 0))
    boxes = [
        ((70, 240, 350, 460), "Question\nutilisateur", PALE_ORANGE, ORANGE),
        ((450, 240, 760, 460), "Détection de la demande\net de la région", PALE_GREEN, GREEN),
        ((860, 240, 1160, 460), "RAG ChromaDB\nsi information\nréglementaire", (242, 247, 255), (72, 119, 180)),
        ((1260, 240, 1540, 460), "Gemini 2.0 Flash\nréponse concise\nen français", PALE_GREEN, GREEN),
    ]
    for box, label, fill, outline in boxes:
        rounded_box(draw, box, label, fill, outline)
    for index in range(len(boxes) - 1):
        arrow(draw, (boxes[index][0][2], 350), (boxes[index + 1][0][0], 350))
    draw.rounded_rectangle((470, 555, 1330, 690), radius=20, outline=ORANGE, width=3, fill=(255, 251, 243))
    draw_centered(draw, (490, 560, 1310, 685), "Réponse ou action guidée : simulation, rendez-vous, modification ou annulation", font(31, True), (120, 79, 15))
    draw.text((120, 708), "Les actions sensibles sont proposées puis exécutées seulement après confirmation explicite de l'utilisateur.", font=font(25), fill=GRAY)
    image.save(path)


def simulation_figure(path: Path):
    image = Image.new("RGB", (1800, 780), "white")
    draw = ImageDraw.Draw(image)
    draw.text((80, 42), "Chaîne de traitement d'une simulation énergétique", font=font(44, True), fill=(0, 0, 0))
    boxes = [
        ((60, 245, 335, 475), "Formulaire\nsurface, consommation,\nlogement, région", PALE_ORANGE, ORANGE),
        ((425, 245, 735, 475), "Validation\ndes informations\nénergétiques", PALE_GREEN, GREEN),
        ((825, 245, 1135, 475), "Calculs métier\néconomies, PEB,\nprimes, score", (242, 247, 255), (72, 119, 180)),
        ((1225, 245, 1535, 475), "Résultat\nrecommandations\net sauvegarde", PALE_GREEN, GREEN),
    ]
    for box, label, fill, outline in boxes:
        rounded_box(draw, box, label, fill, outline)
    for i in range(len(boxes) - 1):
        arrow(draw, (boxes[i][0][2], 360), (boxes[i + 1][0][0], 360))
    draw.rounded_rectangle((550, 570, 1260, 700), radius=20, outline=GREEN, width=3, fill=(243, 250, 245))
    draw_centered(draw, (570, 578, 1240, 692), "Gemini ajoute un conseil textuel uniquement si la clé API est configurée ; les chiffres proviennent des règles métier.", font(27, True), (20, 87, 54))
    image.save(path)


def tests_chart(path: Path):
    image = Image.new("RGB", (1500, 680), "white")
    draw = ImageDraw.Draw(image)
    draw.text((70, 40), "Résultats des tests fonctionnels du chatbot", font=font(42, True), fill=(0, 0, 0))
    draw.text((70, 102), "6 scénarios automatisés validés dans le projet", font=font(27), fill=GRAY)
    draw.rounded_rectangle((220, 210, 800, 420), radius=20, fill=GREEN)
    draw.rounded_rectangle((850, 210, 1280, 420), radius=20, fill=(222, 227, 223), outline=(160, 160, 160), width=2)
    draw_centered(draw, (240, 225, 780, 405), "6\nscénarios validés", font(45, True), (255, 255, 255), spacing=10)
    draw_centered(draw, (870, 225, 1260, 405), "0\néchec observé", font(40, True), (55, 55, 55), spacing=10)
    draw.text((115, 555), "Ces chiffres mesurent les parcours fonctionnels testés ; ils ne constituent pas une évaluation générale du modèle de langage.", font=font(24), fill=GRAY)
    image.save(path)


def corpus_chart(path: Path):
    image = Image.new("RGB", (1600, 760), "white")
    draw = ImageDraw.Draw(image)
    draw.text((78, 40), "Corpus documentaire local inventorié", font=font(44, True), fill=(0, 0, 0))
    draw.text((80, 104), "14 fichiers dans le dossier documents du service IA", font=font(28), fill=GRAY)
    data = [("PDF", 6, GREEN), ("DOCX", 5, (37, 145, 83)), ("TXT", 2, ORANGE), ("CSV", 1, (235, 89, 57))]
    base_y, left, bar_w, gap, unit = 610, 200, 190, 135, 62
    draw.line((130, base_y, 1470, base_y), fill=(150, 150, 150), width=3)
    for value in range(0, 7):
        y = base_y - value * unit
        draw.line((130, y, 1470, y), fill=(232, 232, 232), width=1)
        draw.text((80, y - 14), str(value), font=font(20), fill=GRAY)
    for index, (label, value, color) in enumerate(data):
        x = left + index * (bar_w + gap)
        top = base_y - value * unit
        draw.rounded_rectangle((x, top, x + bar_w, base_y), radius=12, fill=color)
        number = str(value)
        width = draw.textbbox((0, 0), number, font=font(34, True))[2]
        draw.text((x + (bar_w - width) / 2, top - 48), number, font=font(34, True), fill=(0, 0, 0))
        label_width = draw.textbbox((0, 0), label, font=font(28, True))[2]
        draw.text((x + (bar_w - label_width) / 2, 645), label, font=font(28, True), fill=(0, 0, 0))
    draw.text((80, 700), "Statistique descriptive réelle ; elle mesure la composition du corpus et non la performance du modèle.", font=font(24), fill=GRAY)
    image.save(path)


def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), "D9D9D9")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_run_font(run, size=12, bold=False, italic=False, color=BLACK):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def format_paragraph(paragraph, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=6, line=1.12):
    paragraph.alignment = alignment
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="Normal")
    format_paragraph(p)
    if bold_lead:
        lead = p.add_run(bold_lead)
        set_run_font(lead, 12, bold=True)
    run = p.add_run(text)
    set_run_font(run, 12)
    return p


def add_heading(doc, text, level=2, page_break_before=False):
    p = doc.add_paragraph(style="Normal")
    format_paragraph(p, alignment=WD_ALIGN_PARAGRAPH.LEFT, before=10 if level == 2 else 6, after=5, line=1.0)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = page_break_before
    run = p.add_run(text)
    set_run_font(run, 14 if level == 2 else 12, bold=True)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph(style="Normal")
    format_paragraph(p, alignment=WD_ALIGN_PARAGRAPH.CENTER, before=2, after=9, line=1.0)
    run = p.add_run(text)
    set_run_font(run, 10, italic=True)
    return p


def add_table_caption(doc, text):
    p = doc.add_paragraph(style="Normal")
    format_paragraph(p, alignment=WD_ALIGN_PARAGRAPH.CENTER, before=5, after=3, line=1.0)
    run = p.add_run(text)
    set_run_font(run, 10, italic=True)
    return p


def add_table(doc, headers, rows, widths, caption=None):
    if caption:
        add_table_caption(doc, caption)
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, label in enumerate(headers):
        cell = header.cells[index]
        cell.width = Cm(widths[index])
        set_cell_shading(cell, "1A6D45")
        set_cell_margins(cell)
        set_cell_border(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(label)
        set_run_font(run, 10, bold=True, color=RGBColor(255, 255, 255))
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cell = cells[index]
            cell.width = Cm(widths[index])
            set_cell_margins(cell)
            set_cell_border(cell)
            if row_index % 2:
                set_cell_shading(cell, "F4FAF6")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            format_paragraph(p, alignment=WD_ALIGN_PARAGRAPH.LEFT, after=0, line=1.0)
            run = p.add_run(value)
            set_run_font(run, 9)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_figure(doc, image_path: Path, width_cm: float, caption: str):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))
    add_caption(doc, caption)


def remove_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def main():
    WORK.mkdir(exist_ok=True)
    arch = WORK / "figure_architecture_ia.png"
    chatbot = WORK / "figure_chatbot_luna.png"
    corpus = WORK / "figure_corpus.png"
    simulation = WORK / "figure_simulation_pipeline.png"
    tests = WORK / "figure_chatbot_tests.png"
    architecture_figure(arch)
    chatbot_figure(chatbot)
    corpus_chart(corpus)
    simulation_figure(simulation)
    tests_chart(tests)

    document = Document(SOURCE)
    chapter_index = next(
        index for index, paragraph in enumerate(document.paragraphs)
        if "Réalisation du module d’Intelligence Artificielle" in paragraph.text
    )
    chapter_title = document.paragraphs[chapter_index]
    for paragraph in list(document.paragraphs[chapter_index + 1:]):
        remove_paragraph(paragraph)
    chapter_title.text = "4 Réalisation du module d'Intelligence Artificielle"
    format_paragraph(chapter_title, alignment=WD_ALIGN_PARAGRAPH.CENTER, before=8, after=12, line=1.0)
    chapter_title.paragraph_format.page_break_before = True
    chapter_title.paragraph_format.keep_with_next = True
    title_run = chapter_title.runs[0]
    set_run_font(title_run, 16, bold=True)

    add_heading(document, "4.1 Présentation du module IA")
    add_body(document, "Le module d'intelligence artificielle de la plateforme EcoReno+ améliore l'accompagnement de l'utilisateur pendant son projet de rénovation énergétique. Il réunit deux services complémentaires : le chatbot Luna, chargé de la conversation et des actions guidées, et le service IA de simulation, chargé de la recherche documentaire, des calculs et des recommandations.")
    add_body(document, "Cette partie s'appuie sur les services effectivement présents dans le projet. Le chatbot est développé dans un microservice FastAPI distinct ; il expose la route POST /chat et s'appuie sur Gemini 2.0 Flash. Le service IA de simulation et de recherche RAG expose notamment POST /query-rag et communique avec Spring Boot pour traiter les simulations.")

    add_heading(document, "4.2 Architecture IA Python et FastAPI")
    add_body(document, "L'architecture adopte une séparation en microservices. Le frontend EcoReno+ présente les formulaires et les résultats. Spring Boot conserve les données métier, l'authentification, les conversations, les rendez-vous et les simulations. Les deux services FastAPI concentrent les traitements IA. Cette organisation permet de faire évoluer le chatbot, la base de connaissances et le moteur de simulation sans modifier l'ensemble du backend.")
    add_figure(document, arch, 15.8, "Figure Architecture générale du module IA EcoReno+")
    add_table(document, ["Composant", "Responsabilité principale", "Technologie"], [
        ("Frontend EcoReno+", "Saisie des données, affichage du chat et des résultats.", "React"),
        ("Backend métier", "Utilisateurs, formulaires, simulations, messages et rendez-vous.", "Spring Boot"),
        ("Chatbot Luna", "Conversation, actions guidées et confirmations.", "FastAPI et Gemini"),
        ("Service IA", "RAG, règles de simulation, conseils et worker.", "FastAPI et ChromaDB"),
    ], [3.1, 8.1, 5.2], "Table Composants de l'architecture IA")

    add_heading(document, "4.3 Chatbot intelligent Luna")
    add_heading(document, "4.3.1 Objectif et fonctionnalités", level=3)
    add_body(document, "Le chatbot Luna a été développé pour répondre aux questions liées à la rénovation énergétique et guider l'utilisateur dans son parcours. Il peut expliquer les solutions de rénovation, demander la région si elle est nécessaire, orienter vers une simulation, puis aider à créer, modifier ou annuler un rendez-vous. Les messages et le contexte de conversation sont sauvegardés par le backend Spring Boot.")
    add_body(document, "Le chatbot ne déclenche pas directement une action sensible. Avant de créer une simulation ou un rendez-vous, il présente une proposition et demande une confirmation explicite. Si les informations nécessaires au calcul ne sont pas présentes, il redirige l'utilisateur vers le formulaire concerné. Cette logique protège l'utilisateur contre les créations involontaires et améliore la cohérence des données enregistrées.")
    add_figure(document, chatbot, 15.8, "Figure Fonctionnement conversationnel du chatbot Luna")

    add_heading(document, "4.3.2 Choix de l'approche RAG", level=3)
    add_body(document, "Un modèle de langage seul peut produire une réponse naturelle, mais ne garantit pas qu'elle corresponde aux aides, aux réglementations ou aux informations propres à une région. Le projet a donc retenu une approche RAG : avant de répondre à une question réglementaire ou régionale, le chatbot recherche des passages pertinents dans la base documentaire puis les utilise comme contexte de réponse.")
    add_table(document, ["Approche", "Avantage", "Limite dans EcoReno+", "Décision"], [
        ("Règles fixes", "Prévisibles pour des scénarios simples.", "Réponses limitées et difficilement extensibles.", "Non retenue seule"),
        ("LLM seul", "Réponse naturelle rapide.", "Peut manquer de source locale ou métier.", "Non retenu seul"),
        ("Fine tuning", "Spécialisation possible du modèle.", "Nécessite un corpus annoté et un protocole d'entraînement.", "Non implémenté"),
        ("RAG avec LLM", "Réponse fondée sur les documents du projet.", "Dépend de la qualité du corpus et de la recherche.", "Retenu"),
    ], [2.7, 4.0, 6.0, 3.7], "Table Comparaison des approches conversationnelles")
    add_body(document, "Le RAG est plus adapté à cette version du projet que le fine tuning. Il permet d'ajouter ou de mettre à jour une source sans entraîner un nouveau modèle. Il offre également une meilleure traçabilité : les réponses reposent sur des extraits identifiés dans la base de connaissances.")

    add_heading(document, "4.3.3 Corpus, embeddings et recherche", level=3)
    add_body(document, "La base de connaissances contient 14 fichiers locaux : 6 PDF, 5 documents Word, 2 fichiers texte et 1 fichier CSV. Lors de l'ingestion, les documents sont découpés en segments d'environ 800 caractères avec un recouvrement de 100 caractères. Chaque segment peut contenir son document source et sa région, ce qui permet de privilégier les informations pertinentes pour la Wallonie, Bruxelles ou la Flandre.")
    add_body(document, "Les vecteurs sont stockés dans ChromaDB. Le service utilise LocalHashEmbeddings, une représentation locale déterministe de 384 dimensions. Ce choix ne nécessite ni téléchargement de modèle ni clé externe pour la recherche. Il convient à une première version locale et testable ; une évolution possible sera l'adoption d'un modèle sémantique spécialisé lorsque des mesures de qualité de recherche seront disponibles.")
    add_figure(document, corpus, 11.5, "Figure Répartition du corpus documentaire local")

    add_heading(document, "4.3.4 Choix du modèle génératif", level=3)
    add_body(document, "Le modèle actif du chatbot est Gemini 2.0 Flash. Il a été retenu pour produire des réponses courtes et compréhensibles en français, tout en étant encadré par les extraits RAG et les règles métier du projet. La génération est volontairement limitée : le chatbot ne doit pas annoncer une aide, un prix ou une économie chiffrée sans source appropriée.")
    add_table(document, ["Critère", "Gemini 2.0 Flash avec RAG", "LLM seul", "Fine tuning"], [
        ("Réponse conversationnelle", "Oui", "Oui", "Oui"),
        ("Appui sur les sources du projet", "Oui", "Non garanti", "Possible après entraînement"),
        ("Besoin d'un dataset annoté", "Non", "Non", "Oui"),
        ("Mise à jour des connaissances", "Ajout de documents", "Dépend du modèle", "Nouvel entraînement"),
        ("Adaptation à la version actuelle", "Très adaptée", "Partielle", "À envisager plus tard"),
    ], [3.7, 4.3, 3.7, 4.7], "Table Justification du choix du modèle conversationnel")

    add_heading(document, "4.4 Simulation énergétique")
    add_heading(document, "4.4.1 Objectif et données utilisées", level=3)
    add_body(document, "La simulation énergétique estime l'effet potentiel des travaux envisagés à partir des informations saisies dans le formulaire. Les données utiles incluent notamment la surface, la consommation annuelle, le type de logement, le chauffage, l'isolation, la période de construction, l'objectif de rénovation, la région et le produit sélectionné. Le système peut également exploiter les informations PEB disponibles dans son contexte de calcul.")
    add_body(document, "Le résultat rassemble des estimations de consommation ou d'économies, un score, des critères d'analyse, des recommandations et une information indicative sur les primes. Les calculs chiffrés restent séparés des explications rédigées par l'IA ; cette séparation est essentielle pour conserver la traçabilité et l'explicabilité du résultat.")
    add_figure(document, simulation, 15.8, "Figure Chaîne de traitement de la simulation énergétique")

    add_heading(document, "4.4.2 Choix de la méthode de simulation", level=3)
    add_body(document, "La version actuelle ne repose pas sur un modèle de Machine Learning entraîné. Elle adopte une approche hybride : des règles de calcul métier donnent les estimations chiffrées, des informations régionales et PEB enrichissent le contexte, puis Gemini peut produire un conseil textuel si la clé API est configurée. Cette démarche est cohérente avec les données réellement disponibles dans le projet.")
    add_table(document, ["Approche", "Principe", "État réel dans le projet", "Choix"], [
        ("Calcul déterministe", "Règles et fourchettes explicables.", "Implémenté pour économies, primes et comparaison PEB.", "Base du résultat"),
        ("ML supervisé", "Apprentissage sur des historiques de consommation.", "Aucun dataset labellisé ni modèle entraîné trouvé.", "Non déployé"),
        ("Approche hybride", "Règles contrôlées et conseil génératif.", "Règles actives ; conseil Gemini conditionnel.", "Retenue"),
    ], [3.3, 4.9, 5.6, 2.6], "Table Comparaison des approches de simulation")
    add_body(document, "Une régression linéaire, une forêt aléatoire ou XGBoost ne peuvent pas être présentés comme modèles réellement testés dans cette version, car le projet ne contient ni dataset d'entraînement, ni fichier de modèle, ni résultats expérimentaux. Il serait donc incorrect de fabriquer des valeurs MAE, RMSE ou R². Lorsque des données réelles seront collectées, ces modèles pourront être comparés dans un protocole expérimental séparé.")
    add_table(document, ["Métrique pour une future évaluation", "Valeur actuelle", "Condition nécessaire"], [
        ("MAE", "Non disponible", "Consommations ou économies réelles de référence."),
        ("RMSE", "Non disponible", "Jeu de test séparé et prédictions comparables."),
        ("R²", "Non disponible", "Modèle supervisé entraîné sur des données fiables."),
    ], [4.9, 3.6, 7.9], "Table Indicateurs prévus pour une future évaluation ML")

    add_heading(document, "4.4.3 Conseils générés et résultats de test", level=3)
    add_body(document, "Le service de simulation peut appeler le modèle Gemini configuré dans le projet pour transformer les résultats en recommandations simples. Ce modèle ne remplace pas les calculs. Les valeurs chiffrées sont établies par les règles de simulation ; Gemini intervient seulement pour faciliter la compréhension des priorités de rénovation.")
    add_body(document, "Dans un scénario d'intégration enregistré, une simulation a atteint l'état terminée avec un score de 8,0 et quatre recommandations. Cette observation confirme le passage complet entre la collecte des données, le traitement et la restitution. Elle ne constitue pas une mesure générale de précision énergétique ou une validation scientifique de toutes les estimations.")

    add_heading(document, "4.5 Intégration Spring Boot et FastAPI")
    add_body(document, "Spring Boot est le point central des données métier. Le chatbot utilise ses API pour enregistrer les messages, obtenir le contexte de conversation et créer ou modifier les rendez-vous après confirmation. Le service de simulation récupère les simulations en attente, demande leur contexte, puis renvoie le résultat calculé afin qu'il soit sauvegardé et affiché dans le frontend.")
    add_body(document, "Cette intégration évite de laisser le frontend effectuer seul les calculs ou conserver les résultats. Elle permet également de contrôler les droits d'accès, de préserver l'historique de la conversation et de centraliser les données de simulation dans le backend.")

    add_heading(document, "4.6 Résultats et synthèse")
    add_body(document, "Les six scénarios automatisés du chatbot sont validés : confirmation avant simulation, prise en compte du produit et du formulaire, demande de région lorsque nécessaire, redirection en cas de profil incomplet, création de rendez-vous après confirmation et proposition d'un autre créneau en cas de conflit. Ces résultats valident les parcours fonctionnels ciblés ; ils ne représentent pas une mesure universelle de qualité linguistique du modèle.")
    add_figure(document, tests, 13.5, "Figure Résultats des tests fonctionnels du chatbot")
    add_body(document, "Le module IA d'EcoReno+ est donc fondé sur des éléments réellement implémentés : chatbot Luna connecté au backend, Gemini 2.0 Flash encadré par des règles métier, RAG sur ChromaDB, corpus local de 14 documents et simulation hybride explicable. L'évolution la plus pertinente consiste à enrichir le corpus documentaire et à constituer un jeu de données de référence pour comparer, avec de vraies métriques, des modèles prédictifs de consommation ou d'économies.")

    for section in document.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(1.0)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    document.core_properties.title = "Rapport PFA EcoReno Chapitre Intelligence Artificielle"
    document.core_properties.subject = "Réalisation du module d'Intelligence Artificielle"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
